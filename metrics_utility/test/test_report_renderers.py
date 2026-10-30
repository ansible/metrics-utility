"""Behavioral checks for shared report definitions and spreadsheet output."""

import datetime
import gc
import json
import weakref

from pathlib import Path
from unittest.mock import Mock
from zipfile import ZipFile

import numpy as np
import pandas as pd
import pytest

from defusedxml import ElementTree
from openpyxl import load_workbook

from metrics_utility.automation_controller_billing.report.definition import (
    Alignment,
    Border,
    CellReference,
    Font,
    Multiply,
    PatternFill,
    Side,
    Sum,
    Workbook,
)
from metrics_utility.automation_controller_billing.report.renderers.ods import render as render_ods
from metrics_utility.automation_controller_billing.report.renderers.xlsx import render as render_xlsx
from metrics_utility.automation_controller_billing.report_saver.report_saver_directory import ReportSaverDirectory
from metrics_utility.automation_controller_billing.report_saver.report_saver_s3 import ReportSaverS3
from metrics_utility.test.util import run_build_int


TABLE = 'urn:oasis:names:tc:opendocument:xmlns:table:1.0'
OFFICE = 'urn:oasis:names:tc:opendocument:xmlns:office:1.0'
NS = {'table': TABLE, 'office': OFFICE}


def _definition():
    book = Workbook()
    book.remove(book.active)
    sheet = book.create_sheet('Usage')
    sheet.column_dimensions[1].width = 24
    sheet.row_dimensions[1].height = 25
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
    sheet.cell(1, 1).value = 'Usage report'
    sheet.cell(1, 1).font = Font(name='Arial', size=12, color='000000', bold=True)
    sheet.cell(2, 1).value = 3
    sheet.cell(2, 2).value = 12.5
    sheet.cell(2, 3).value = Multiply(CellReference(2, 1), CellReference(2, 2))
    sheet.cell(3, 3).value = Sum(CellReference(2, 3), CellReference(2, 3))
    sheet.cell(2, 3).number_format = '$#,##0.00'
    sheet.cell(2, 3).border = Border(left=Side(border_style='dotted', color='000000'))
    status = book.create_sheet('Collection status')
    status.cell(1, 1).value = 'Warning\nSecond line'
    status.cell(1, 1).fill = PatternFill('solid', fg_color='ffcc17')
    status.cell(1, 1).alignment = Alignment(horizontal='center', vertical='center')
    status.cell(2, 1).value = datetime.timedelta(hours=36)
    status.cell(3, 1).value = datetime.datetime(2024, 4, 17, 12, 19, 38, 762500)
    status.cell(4, 1).value = pd.NaT
    status.cell(5, 1).value = datetime.timedelta(hours=36, microseconds=123456)
    status.cell(6, 1).value = np.bool_(True)
    status.cell(7, 1).value = datetime.date(2024, 4, 17)
    status.cell(8, 1).value = datetime.time(12, 34, 56, 123456)
    status.cell(9, 1).fill = PatternFill('solid', fg_color='92d050')
    status.cell(10, 1).fill = PatternFill('solid', fg_color='FF0000')
    return book


def test_shared_definition_renders_formulas_styles_and_typed_values(tmp_path):
    definition = _definition()
    xlsx_path = tmp_path / 'report.xlsx'
    ods_path = tmp_path / 'report.ods'
    render_xlsx(definition).save(xlsx_path)
    render_ods(definition).save(ods_path)

    xlsx = load_workbook(xlsx_path)
    assert xlsx.sheetnames == ['Usage', 'Collection status']
    assert xlsx['Usage']['C2'].value == '=A2*B2'
    assert xlsx['Usage']['C3'].value == '=SUM(C2:C2)'
    assert 'A1:B1' in {str(item) for item in xlsx['Usage'].merged_cells.ranges}
    assert xlsx['Collection status']['A1'].fill.fgColor.rgb[-6:] == 'ffcc17'
    assert xlsx['Collection status']['A2'].value == datetime.timedelta(hours=36)

    with ZipFile(ods_path) as archive:
        assert archive.read('mimetype') == b'application/vnd.oasis.opendocument.spreadsheet'
        document = ElementTree.fromstring(archive.read('content.xml'))
        sheets = document.findall('.//table:table', NS)
        assert [sheet.attrib[f'{{{TABLE}}}name'] for sheet in sheets] == ['Usage', 'Collection status']
        cells = document.findall('.//table:table-cell', NS)
        formulas = [cell.attrib[f'{{{TABLE}}}formula'] for cell in cells if f'{{{TABLE}}}formula' in cell.attrib]
        assert formulas == ['of:=[.A2]*[.B2]', 'of:=SUM([.C2:.C2])']
        assert any(cell.attrib.get(f'{{{OFFICE}}}time-value') == 'PT129600S' for cell in cells)
        assert any(cell.attrib.get(f'{{{OFFICE}}}time-value') == 'PT129600.123456S' for cell in cells)
        assert any(cell.attrib.get(f'{{{OFFICE}}}boolean-value') == 'true' for cell in cells)
        assert any(cell.attrib.get(f'{{{OFFICE}}}date-value') == '2024-04-17T12:19:38.762500' for cell in cells)
        assert any(cell.attrib.get(f'{{{OFFICE}}}date-value') == '2024-04-17' for cell in cells)
        assert any(cell.attrib.get(f'{{{OFFICE}}}time-value') == 'PT45296.123456S' for cell in cells)
        assert any(cell.attrib.get(f'{{{TABLE}}}number-columns-spanned') == '2' for cell in cells)
        styles = archive.read('content.xml')
        assert b'#ffcc17' in styles
        assert b'#92d050' in styles
        assert b'#FF0000' in styles


def test_ods_saver_writes_exact_s3_temporary_path(tmp_path, monkeypatch):
    uploaded = {}

    def upload(path, destination):
        uploaded['destination'] = destination
        uploaded['bytes'] = Path(path).read_bytes()

    handler = Mock()
    handler.upload_file.side_effect = upload
    monkeypatch.setattr('metrics_utility.automation_controller_billing.report_saver.report_saver_s3.S3Handler', lambda **_: handler)
    saver = ReportSaverS3({'report_spreadsheet_destination_path': 'reports/sample.ods'})
    saver.save(render_ods(_definition()))
    assert uploaded['destination'] == 'reports/sample.ods'
    from io import BytesIO

    with ZipFile(BytesIO(uploaded['bytes'])) as archive:
        assert archive.read('mimetype') == b'application/vnd.oasis.opendocument.spreadsheet'


def test_ods_rejects_report_without_sheets():
    definition = Workbook()
    definition.remove(definition.active)
    with pytest.raises(ValueError, match='at least one sheet'):
        render_ods(definition)


def test_ods_formula_values_are_cached_and_readable(tmp_path):
    path = tmp_path / 'report.ods'
    render_ods(_definition()).save(path)

    with ZipFile(path) as archive:
        document = ElementTree.fromstring(archive.read('content.xml'))
    cells = document.findall('.//table:table-cell', NS)
    formula_cells = {cell.attrib[f'{{{TABLE}}}formula']: cell for cell in cells if f'{{{TABLE}}}formula' in cell.attrib}
    assert float(formula_cells['of:=[.A2]*[.B2]'].attrib[f'{{{OFFICE}}}value']) == 37.5
    assert float(formula_cells['of:=SUM([.C2:.C2])'].attrib[f'{{{OFFICE}}}value']) == 37.5
    values = pd.read_excel(path, sheet_name='Usage', header=None)
    assert values.iloc[1, 2] == 37.5
    assert values.iloc[2, 2] == 37.5


def test_ods_formula_cache_handles_empty_cells_and_unavailable_results(tmp_path):
    definition = Workbook()
    sheet = definition.active
    sheet.cell(1, 1).value = Sum(CellReference(2, 3), CellReference(4, 3))
    sheet.cell(2, 1).value = 3
    sheet.cell(2, 2).value = 12.5
    sheet.cell(2, 3).value = Multiply(CellReference(2, 1), CellReference(2, 2))
    sheet.cell(3, 3).value = Multiply(CellReference(3, 1), CellReference(3, 2))
    sheet.cell(4, 3).value = -2
    sheet.cell(5, 1).value = 'not a number'
    sheet.cell(5, 3).value = Multiply(CellReference(5, 1), CellReference(2, 1))
    sheet.cell(6, 3).value = Sum(CellReference(5, 3), CellReference(5, 3))
    sheet.cell(7, 3).value = Multiply(CellReference(7, 3), CellReference(2, 1))
    path = tmp_path / 'report.ods'
    render_ods(definition).save(path)

    with ZipFile(path) as archive:
        document = ElementTree.fromstring(archive.read('content.xml'))
    cells = document.findall('.//table:table-cell', NS)
    formulas = {cell.attrib[f'{{{TABLE}}}formula']: cell.attrib for cell in cells if f'{{{TABLE}}}formula' in cell.attrib}
    assert float(formulas['of:=SUM([.C2:.C4])'][f'{{{OFFICE}}}value']) == 35.5
    assert float(formulas['of:=[.A3]*[.B3]'][f'{{{OFFICE}}}value']) == 0
    assert f'{{{OFFICE}}}value' not in formulas['of:=[.A5]*[.A2]']
    assert f'{{{OFFICE}}}value' not in formulas['of:=SUM([.C5:.C5])']
    assert f'{{{OFFICE}}}value' not in formulas['of:=[.C7]*[.A2]']


@pytest.mark.parametrize(
    ('value', 'expected'),
    [
        ('First\r\nSecond', 'First\nSecond'),
        ('First\rSecond', 'First\nSecond'),
        ('First\nSecond', 'First\nSecond'),
        ('First\r\n\r\nSecond', 'First\n\nSecond'),
    ],
)
def test_ods_preserves_text_line_breaks(tmp_path, value, expected):
    definition = Workbook()
    definition.active.cell(1, 1).value = value
    path = tmp_path / 'report.ods'
    render_ods(definition).save(path)

    with ZipFile(path) as archive:
        document = ElementTree.fromstring(archive.read('content.xml'))
    text_namespace = 'urn:oasis:names:tc:opendocument:xmlns:text:1.0'
    paragraph = document.find(f'.//{{{TABLE}}}table-cell/{{{text_namespace}}}p')
    parts = [paragraph.text or '']
    for child in paragraph:
        if child.tag == f'{{{text_namespace}}}line-break':
            parts.append('\n')
        parts.append(child.tail or '')
    assert ''.join(parts) == expected


def test_empty_sheet_title_has_same_default_name_in_both_formats(tmp_path):
    definition = Workbook()
    definition.create_sheet('')
    xlsx_path, ods_path = tmp_path / 'report.xlsx', tmp_path / 'report.ods'
    render_xlsx(definition).save(xlsx_path)
    render_ods(definition).save(ods_path)

    assert load_workbook(xlsx_path).sheetnames == ['Sheet', 'Sheet1']
    with ZipFile(ods_path) as archive:
        document = ElementTree.fromstring(archive.read('content.xml'))
    names = [sheet.attrib[f'{{{TABLE}}}name'] for sheet in document.findall('.//table:table', NS)]
    assert names == ['Sheet', 'Sheet1']


def test_build_report_releases_ods_definition_before_saving(tmp_path, monkeypatch):
    root = Path(__file__).resolve().parents[2]
    (tmp_path / 'data').symlink_to(root / 'metrics_utility/test/test_data/data', target_is_directory=True)
    snapshot = root / 'metrics_utility/test/snapshot_tests/CCSP/data/CCSPv2/snapshot_def_2024-04.json'
    env = json.loads(snapshot.read_text())['env_vars']
    env['METRICS_UTILITY_SHIP_PATH'] = str(tmp_path)
    captured = {}
    original_save = ReportSaverDirectory.save

    def track_definition(definition):
        captured['definition'] = weakref.ref(definition)
        return render_ods(definition)

    def verify_before_save(saver, spreadsheet):
        gc.collect()
        assert captured['definition']() is None
        original_save(saver, spreadsheet)

    monkeypatch.setattr('metrics_utility.automation_controller_billing.report.renderers.ods.render', track_definition)
    monkeypatch.setattr(ReportSaverDirectory, 'save', verify_before_save)
    run_build_int(env, {'month': '2024-04', 'format': 'ods'})
    assert (tmp_path / 'reports/2024/04/CCSPv2-2024-04.ods').is_file()


@pytest.mark.parametrize('report_type', ['CCSP', 'CCSPv2'])
def test_build_report_writes_xlsx_and_ods_to_separate_paths(tmp_path, report_type, monkeypatch):
    root = Path(__file__).resolve().parents[2]
    input_data = root / 'metrics_utility/test/test_data/data'
    (tmp_path / 'data').symlink_to(input_data, target_is_directory=True)
    snapshot = root / f'metrics_utility/test/snapshot_tests/CCSP/data/{report_type}/snapshot_def_2024-04.json'
    env = json.loads(snapshot.read_text())['env_vars']
    env['METRICS_UTILITY_SHIP_PATH'] = str(tmp_path)
    run_build_int(env, {'month': '2024-04'})
    run_build_int(env, {'month': '2024-04', 'format': 'ods'})

    report_dir = tmp_path / 'reports/2024/04'
    xlsx = report_dir / f'{report_type}-2024-04.xlsx'
    ods = report_dir / f'{report_type}-2024-04.ods'
    assert xlsx.is_file()
    assert ods.is_file()
    assert load_workbook(xlsx).sheetnames[0] == 'Usage Reporting'
    with ZipFile(ods) as archive:
        document = ElementTree.fromstring(archive.read('content.xml'))
        assert document.find('.//table:table', NS).attrib[f'{{{TABLE}}}name'] == 'Usage Reporting'
    with monkeypatch.context() as patch:
        patch.setattr(
            'metrics_utility.automation_controller_billing.report.renderers.ods.render',
            Mock(side_effect=AssertionError('existing report must be skipped')),
        )
        run_build_int(env, {'month': '2024-04', 'format': 'ods'})
    ods.write_bytes(b'old report')
    run_build_int(env, {'month': '2024-04', 'format': 'ods', 'force': True})
    with ZipFile(ods) as archive:
        assert archive.read('mimetype') == b'application/vnd.oasis.opendocument.spreadsheet'


@pytest.mark.parametrize('report_type', ['CCSP', 'CCSPv2'])
def test_ccsp_ods_formula_results_are_readable_with_pandas(tmp_path, report_type):
    root = Path(__file__).resolve().parents[2]
    (tmp_path / 'data').symlink_to(root / 'metrics_utility/test/test_data/data', target_is_directory=True)
    snapshot = root / f'metrics_utility/test/snapshot_tests/CCSP/data/{report_type}/snapshot_def_2024-04.json'
    env = json.loads(snapshot.read_text())['env_vars']
    env['METRICS_UTILITY_SHIP_PATH'] = str(tmp_path)
    env['METRICS_UTILITY_PRICE_PER_NODE'] = '12.5'
    run_build_int(env, {'month': '2024-04', 'format': 'ods'})

    path = tmp_path / f'reports/2024/04/{report_type}-2024-04.ods'
    values = pd.read_excel(path, sheet_name='Usage Reporting', header=None)
    if report_type == 'CCSP':
        assert values.iloc[17, 4] == 3212.5
        assert values.iloc[18, 4] == 12.5
        assert values.iloc[19, 4] == 3225
    else:
        assert values.iloc[6, 9] == 3225
        assert values.iloc[1, 9] == 3225
        assert values.iloc[7, 9] == 0


def test_ccspv2_ods_includes_optional_sheets_and_status_colors(tmp_path):
    root = Path(__file__).resolve().parents[2]
    (tmp_path / 'data').symlink_to(root / 'metrics_utility/test/test_data/data', target_is_directory=True)
    env = {
        'METRICS_UTILITY_SHIP_PATH': str(tmp_path),
        'METRICS_UTILITY_SHIP_TARGET': 'directory',
        'METRICS_UTILITY_REPORT_TYPE': 'CCSPv2',
        'METRICS_UTILITY_OPTIONAL_COLLECTORS': 'main_host,main_jobevent,main_indirectmanagednodeaudit',
        'METRICS_UTILITY_OPTIONAL_CCSP_REPORT_SHEETS': (
            'ccsp_summary,jobs,managed_nodes,indirectly_managed_nodes,inventory_scope,usage_by_organizations,'
            'usage_by_collections,usage_by_roles,usage_by_modules,data_collection_status'
        ),
        'METRICS_UTILITY_DEDUPLICATOR': 'ccsp-experimental',
    }
    run_build_int(env, {'since': '2025-07-08', 'until': '2025-07-11', 'format': 'ods'})
    report = tmp_path / 'reports/2025/07/CCSPv2-2025-07-08--2025-07-11.ods'
    with ZipFile(report) as archive:
        content = archive.read('content.xml')
        doc = ElementTree.fromstring(content)
        names = [sheet.attrib[f'{{{TABLE}}}name'] for sheet in doc.findall('.//table:table', NS)]
        assert names == [
            'Usage Reporting',
            'Jobs',
            'Managed nodes',
            'Indirectly Managed nodes',
            'Inventory Scope',
            'Usage by organizations',
            'Usage by collections',
            'Usage by roles',
            'Usage by modules',
            'Data collection status',
        ]
        assert b'#FF0000' in content
        assert b'of:=SUM([.J7:.J12])' in content


def test_renewal_definition_renders_optional_sheets(setup_processed_dataframe, tmp_path):
    from metrics_utility.automation_controller_billing.report.report_renewal_guidance import ReportRenewalGuidance

    options = {
        'ephemeral_days': 30,
        'report_period': '2025-01-01, 2025-06-03',
        'report_type': 'RENEWAL_GUIDANCE',
        'since_date': '2025-01-01',
        'until_date': '2025-06-03',
        'optional_sheets': ['managed_nodes'],
    }
    report = ReportRenewalGuidance({'host_metric': setup_processed_dataframe}, options).build_definition()
    output = tmp_path / 'renewal.ods'
    render_ods(report).save(output)
    with ZipFile(output) as archive:
        doc = ElementTree.fromstring(archive.read('content.xml'))
        names = [sheet.attrib[f'{{{TABLE}}}name'] for sheet in doc.findall('.//table:table', NS)]
        assert names == [
            'Usage Reporting',
            'Managed nodes',
            'Managed nodes ephemeral',
            'Managed nodes ephemeral usage',
            'Deleted Managed nodes',
        ]
