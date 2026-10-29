"""Exercise single-sheet builds through real tarball extraction, not prebuilt dataframes."""

import csv
import io
import json
import tarfile

from contextlib import closing

from openpyxl import load_workbook

from metrics_utility.test.util import run_build_int, utcdt


def _write_collection(ship_path, name, rows):
    """Write one collector's CSV and installation context in an isolated daily archive."""
    directory = ship_path / 'data' / '2025' / '06' / '01'
    directory.mkdir(parents=True, exist_ok=True)
    csv_data = io.StringIO()
    writer = csv.DictWriter(csv_data, fieldnames=list(rows[0]))
    writer.writeheader()
    writer.writerows(rows)

    with tarfile.open(directory / f'fixture-0-{name}.tar.gz', 'w:gz') as archive:
        for filename, content in {
            'config.json': json.dumps({'install_uuid': '11111111-1111-1111-1111-111111111111'}),
            f'{name}.csv': csv_data.getvalue(),
        }.items():
            data = content.encode()
            member = tarfile.TarInfo(filename)
            member.size = len(data)
            archive.addfile(member, io.BytesIO(data))


def _build_single_sheet(ship_path, sheet):
    """Override the whole default sheet list, so unrelated flags cannot enable the input."""
    run_build_int(
        {
            'METRICS_UTILITY_SHIP_TARGET': 'directory',
            'METRICS_UTILITY_SHIP_PATH': str(ship_path),
            'METRICS_UTILITY_REPORT_TYPE': 'CCSPv2',
            'METRICS_UTILITY_OPTIONAL_CCSP_REPORT_SHEETS': sheet,
            'METRICS_UTILITY_DEDUPLICATOR': None,
            'METRICS_UTILITY_ORGANIZATION_FILTER': None,
        },
        {'since': '2025-06-01', 'until': '2025-06-01', 'force': True},
    )
    report = ship_path / 'reports' / '2025' / '06' / 'CCSPv2-2025-06-01--2025-06-01.xlsx'
    assert report.is_file(), f'{sheet} did not produce a workbook despite populated source data'
    return report


def test_jobs_only_loads_summaries_and_renders_job_totals(tmp_path):
    """A jobs-only build must not need ccsp_summary or managed_nodes to load summaries."""
    row = {
        'host_name': 'node-a',
        'ansible_host_variable': '192.0.2.1',
        'organization_name': 'Example org',
        'job_template_name': 'Run automation',
        'job_remote_id': 100,
        'created': '2025-06-01T10:01:00Z',
        'job_created': '2025-06-01T10:00:00Z',
        'dark': 1,
        'failures': 0,
        'ok': 2,
        'skipped': 0,
        'ignored': 0,
        'rescued': 0,
    }
    _write_collection(
        tmp_path,
        'job_host_summary',
        [
            row,
            {**row, 'host_name': 'node-b', 'ansible_host_variable': None, 'dark': 0, 'ok': 3},
            {**row, 'job_remote_id': 101, 'job_created': '2025-06-01T11:00:00Z', 'dark': 0, 'ok': 1},
            {**row, 'host_name': 'unreachable', 'job_remote_id': 102, 'dark': 4, 'ok': 0},
        ],
    )

    with closing(load_workbook(_build_single_sheet(tmp_path, 'jobs'))) as workbook:
        assert workbook.sheetnames == ['Jobs']
        header, data = list(workbook['Jobs'].values)
        assert dict(zip((value.replace('\n', ' ') for value in header), data)) == {
            'Job template name': 'Run automation',
            'Organization name': 'Example org',
            'Job runs': 2,
            'Unique managed nodes automated': 2,
            'Non-unique managed nodes automated': 3,
            'Number of task runs': 7,
            'First run': utcdt('2025-06-01T10:00:00Z').replace(tzinfo=None),
            'Last run': utcdt('2025-06-01T11:00:00Z').replace(tzinfo=None),
        }


def test_infrastructure_summary_only_loads_audits_and_renders_node_counts(tmp_path):
    """An infrastructure-only build must load audits without any managed-node sheet."""
    row = {
        'host_name': 'vm-a',
        'organization_name': 'Example org',
        'job_template_name': 'Discover infrastructure',
        'job_remote_id': 100,
        'created': '2025-06-01T10:01:00Z',
        'job_created': '2025-06-01T10:00:00Z',
        'task_runs': 2,
        'canonical_facts': json.dumps({'asset_id': 'vm-a'}),
        'facts': json.dumps({'infra_type': 'Cloud', 'infra_bucket': 'Public', 'device_type': 'VM'}),
        'events': json.dumps(['example.cloud.discover']),
    }
    _write_collection(
        tmp_path,
        'main_indirectmanagednodeaudit',
        [row, {**row, 'host_name': 'vm-b'}, {**row, 'job_remote_id': 101}],
    )

    with closing(load_workbook(_build_single_sheet(tmp_path, 'infrastructure_summary'))) as workbook:
        assert workbook.sheetnames == ['Infrastructure Summary']
        assert list(workbook['Infrastructure Summary'].values) == [
            ('Infrastructure', 'Device Category', 'Device Type', 'Unique Nodes', 'Total Nodes'),
            ('Cloud', None, None, None, None),
            (None, 'Public', None, None, None),
            (None, None, 'VM', 2, 3),
        ]
