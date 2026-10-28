"""Render a report definition as an openpyxl workbook."""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from metrics_utility.automation_controller_billing.report.definition import Multiply, Sum


def _reference(ref):
    return f'{get_column_letter(ref.column)}{ref.row}'


def _value(value):
    if isinstance(value, Multiply):
        return f'={_reference(value.left)}*{_reference(value.right)}'
    if isinstance(value, Sum):
        return f'=SUM({_reference(value.start)}:{_reference(value.end)})'
    return value


def _border_side(side):
    if side is None:
        return None
    return Side(border_style=side.border_style, color=side.color)


def render(definition):
    workbook = Workbook()
    workbook.remove(workbook.active)
    fonts, fills, borders, alignments = {}, {}, {}, {}
    for sheet in definition.worksheets:
        ws = workbook.create_sheet(sheet.title)
        for start_row, start_column, end_row, end_column in sheet.merges:
            ws.merge_cells(start_row=start_row, start_column=start_column, end_row=end_row, end_column=end_column)
        for source in sheet.cells.values():
            cell = ws.cell(source.row, source.column)
            if source.value is not None:
                cell.value = _value(source.value)
            if source.font is not None:
                if source.font not in fonts:
                    fonts[source.font] = Font(name=source.font.name, size=source.font.size, color=source.font.color, bold=source.font.bold)
                cell.font = fonts[source.font]
            if source.fill is not None:
                if source.fill not in fills:
                    fills[source.fill] = PatternFill(source.fill.pattern, fgColor=source.fill.fgColor)
                cell.fill = fills[source.fill]
            if source.border is not None:
                if source.border not in borders:
                    borders[source.border] = Border(
                        left=_border_side(source.border.left),
                        right=_border_side(source.border.right),
                        top=_border_side(source.border.top),
                        bottom=_border_side(source.border.bottom),
                    )
                cell.border = borders[source.border]
            if source.alignment is not None:
                if source.alignment not in alignments:
                    alignments[source.alignment] = Alignment(horizontal=source.alignment.horizontal, vertical=source.alignment.vertical)
                cell.alignment = alignments[source.alignment]
            if source.number_format is not None:
                cell.number_format = source.number_format
        for column, dimension in sheet.column_dimensions.items():
            if dimension.width is not None:
                ws.column_dimensions[get_column_letter(column)].width = dimension.width
        for row, dimension in sheet.row_dimensions.items():
            if dimension.height is not None:
                ws.row_dimensions[row].height = dimension.height
    return workbook
